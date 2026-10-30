import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from django.conf import settings
from django.core import mail, signing
from django.core.cache import cache
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from .models import NewsletterSubscriber
from .views import NEWSLETTER_SALT


class NewsletterSecurityTests(TestCase):
    def setUp(self):
        cache.clear()
        self.url = reverse('newsletter_subscribe')

    def _post(self, email):
        return self.client.post(self.url, {'email': email, 'next': '/'})

    def test_email_moi_va_email_da_dang_ky_cung_mot_thong_bao(self):
        NewsletterSubscriber.objects.create(email='old@example.com', is_confirmed=True)
        r1 = self._post('new@example.com')
        r2 = self._post('old@example.com')
        self.assertEqual(r1['Location'], r2['Location'])
        self.assertIn('newsletter=pending', r1['Location'])

    def test_chua_xac_nhan_thi_chua_vao_danh_sach(self):
        self._post('new@example.com')
        self.assertFalse(NewsletterSubscriber.objects.filter(email='new@example.com').exists())
        self.assertEqual(len(mail.outbox), 1)

    def test_bam_link_xac_nhan_moi_duoc_dang_ky(self):
        token = signing.dumps('new@example.com', salt=NEWSLETTER_SALT)
        r = self.client.get(reverse('newsletter_confirm', args=[token]))
        self.assertIn('newsletter=confirmed', r['Location'])
        self.assertTrue(NewsletterSubscriber.objects.get(email='new@example.com').is_confirmed)

    def test_token_gia_bi_tu_choi(self):
        r = self.client.get(reverse('newsletter_confirm', args=['token-gia']))
        self.assertIn('newsletter=confirm_invalid', r['Location'])
        self.assertEqual(NewsletterSubscriber.objects.count(), 0)

    def test_gioi_han_so_lan_gui(self):
        for i in range(5):
            self.assertIn('pending', self._post(f'u{i}@example.com')['Location'])
        self.assertIn('newsletter=limited', self._post('u6@example.com')['Location'])


# ---------------------------------------------------------------------------
# SCRUM-36 - Độ dài email đăng ký nhận tin: Phân hoạch tương đương (EP) & Giá trị biên (BVA).
# Input : docs/email_dangky_testcase.xlsx (sheet "Test case")
# Output: docs/email_dangky_report.xlsx
# Chạy: python manage.py test contact -v 2

DOCS_DIR = Path(settings.BASE_DIR) / 'docs'
TESTCASE_FILE = DOCS_DIR / 'email_dangky_testcase.xlsx'
REPORT_FILE = DOCS_DIR / 'email_dangky_report.xlsx'

PENDING_MESSAGE = 'Nếu email hợp lệ, chúng tôi đã gửi thư xác nhận. Vui lòng kiểm tra hộp thư để hoàn tất đăng ký.'
INVALID_MESSAGE = 'Email không hợp lệ, vui lòng kiểm tra lại.'
CONFIRM_LINK = re.compile(r'https?://\S+/newsletter/confirm/\S+/')


def load_test_cases():
    """Đọc sheet "Test case": dòng có ô đầu là "STT" là tiêu đề, các dòng sau là dữ liệu."""
    ws = load_workbook(TESTCASE_FILE, read_only=True)['Test case']
    rows = ws.iter_rows(values_only=True)
    header = next(row for row in rows if row and row[0] == 'STT')
    cases = [dict(zip(header, row)) for row in rows if row and row[1]]
    ws.parent.close()
    return cases


def label_lengths(email):
    domain = email.rsplit('@', 1)[1]
    return ', '.join(str(len(label)) for label in domain.split('.')[:-1])


class NewsletterEmailLengthTests(TestCase):
    """Mỗi dòng trong sheet "Test case" sinh ra một test_TC_NL_xx."""

    results = []

    def setUp(self):
        # Mỗi ca độc lập với giới hạn 5 lần / 10 phút / IP của SCRUM-34
        cache.clear()
        self.client = Client(raise_request_exception=False)

    def run_case(self, case):
        email = case['Email đầu vào']
        # Bảo vệ dữ liệu thiết kế: độ dài ghi trong Excel phải khớp với email thật
        self.assertEqual(len(email), case['Tổng độ dài'], 'Sai dữ liệu: tổng độ dài')
        self.assertEqual(len(email.rsplit('@', 1)[0]), case['Độ dài phần trước @'], 'Sai dữ liệu: phần trước @')
        self.assertEqual(label_lengths(email), str(case['Độ dài nhãn tên miền']), 'Sai dữ liệu: nhãn tên miền')

        started_at = datetime.now()
        start = time.perf_counter()
        response = self.client.post(reverse('newsletter_subscribe'), {'email': email, 'next': '/'}, follow=True)
        duration_ms = (time.perf_counter() - start) * 1000
        final_url = response.redirect_chain[-1][0] if response.redirect_chain else ''
        status = parse_qs(urlparse(final_url).query).get('newsletter', [''])[0]
        content = response.content.decode()
        message = next((m for m in (PENDING_MESSAGE, INVALID_MESSAGE) if m in content), '(không có thông báo)')
        mails = [m for m in mail.outbox if email.lower() in m.to]

        # Bấm link trong thư xác nhận (nếu có) để xem email có thực sự vào DB không
        link = CONFIRM_LINK.search(mails[0].body) if mails else None
        if link:
            self.client.get(urlparse(link.group()).path)
        saved = NewsletterSubscriber.objects.filter(email=email.lower(), is_confirmed=True).exists()

        # Expected / Actual: True = hệ thống chấp nhận email (gửi thư xác nhận, bấm link thì được lưu DB)
        expected = case['Loại'] == 'Hợp lệ'
        accepted = status == 'pending' and message == PENDING_MESSAGE and len(mails) == 1 and saved
        rejected = status == 'invalid' and message == INVALID_MESSAGE and not mails and not saved
        passed = accepted if expected else rejected
        self.results.append({
            **case, 'Expected': expected, 'Actual': accepted, 'Status': 'PASS' if passed else 'FAIL',
            'Thoi_Gian_Chay': started_at.strftime('%d/%m/%Y %H:%M'), 'Thoi_Luong_ms': round(duration_ms, 3),
        })

        actual = (f'?newsletter={status or "(trống)"}, hiển thị "{message}", gửi {len(mails)} thư, '
                  f'{"email ĐƯỢC lưu vào DB" if saved else "không lưu email vào DB"}')
        self.assertTrue(passed, f'{case["Mã TC"]}: mong đợi "{case["Kết quả mong đợi"]}", thực tế {actual}')

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        if cls.results:
            write_report(sorted(cls.results, key=lambda r: r['STT']))


def make_test(case):
    def test(self):
        self.run_case(case)
    test.__doc__ = f'{case["Mã TC"]}: {case["Tên test case"]}'
    return test


for _case in load_test_cases():
    setattr(NewsletterEmailLengthTests, f'test_{_case["Mã TC"].replace("-", "_")}', make_test(_case))


# (tiêu đề cột report, cột tương ứng trong sheet "Test case" hoặc trong kết quả chạy, độ rộng)
REPORT_COLUMNS = [
    ('TC_ID', 'Mã TC', 11),
    ('Ten_Test_Case', 'Tên test case', 36),
    ('Loai', 'Loại', 13),
    ('email', 'Email đầu vào', 40),
    ('tongDoDai', 'Tổng độ dài', 11),
    ('doDaiTruocA', 'Độ dài phần trước @', 13),
    ('doDaiNhan', 'Độ dài nhãn tên miền', 15),
    ('Expected', 'Expected', 10),
    ('Actual', 'Actual', 10),
    ('Status', 'Status', 9),
    ('Tags', 'Tag được bao phủ', 26),
    ('Thoi_Gian_Chay', 'Thoi_Gian_Chay', 18),
    ('Thoi_Luong_ms', 'Thoi_Luong_ms', 15),
]


def write_report(results):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Report'
    ws.append([title for title, _, _ in REPORT_COLUMNS])
    for r in results:
        ws.append([r[key] for _, key, _ in REPORT_COLUMNS])

    for column, (title, _, width) in zip(ws.columns, REPORT_COLUMNS):
        ws.column_dimensions[column[0].column_letter].width = width
        for cell in column:
            cell.alignment = Alignment(horizontal='left' if title == 'email' else 'center')
            if title == 'Thoi_Luong_ms':
                cell.number_format = '0.000'
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'B2'

    wb.save(REPORT_FILE)
